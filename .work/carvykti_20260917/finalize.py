# ai coding: 将基线对比置于首页并核对原始数据和超链接完整性 2026/09/17: 16:42
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
import openpyxl

p=Path('/Users/zendu/Documents/Invest/outputs/carvykti_20260917/Carvykti_基线对比整理.xlsx')
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with ZipFile(p) as z:
    entries={n:z.read(n) for n in z.namelist()}
root=etree.fromstring(entries['xl/workbook.xml'])
sheets=root.find('s:sheets',ns)
old=list(sheets)
order=[x for x in old if x.get('name')=='基线对比']+[x for x in old if x.get('name') not in ('基线对比','基线来源')]+[x for x in old if x.get('name')=='基线来源']
sheets[:]=order
for view in root.findall('s:bookViews/s:workbookView',ns):
    view.set('activeTab','0')
    view.set('firstSheet','0')
entries['xl/workbook.xml']=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(p,'w',ZIP_DEFLATED) as z:
    for name,data in entries.items():z.writestr(name,data)
a=openpyxl.load_workbook('/Users/zendu/Documents/Invest/医药/临床/Carvykti.xlsx')
b=openpyxl.load_workbook(p)
for sheet in a:
    for row in sheet:
        for c in row:
            got=b[sheet.title][c.coordinate]
            assert c.value==got.value,(sheet.title,c.coordinate)
            if c.hyperlink:assert got.hyperlink and got.hyperlink.target==c.hyperlink.target
cached=openpyxl.load_workbook(p,data_only=True)
assert abs(cached['基线对比']['C23'].value-28/286)<1e-10
assert abs(cached['基线对比']['B11'].value-1/208)<1e-10
assert b['基线对比'].freeze_panes=='B2'
assert b.sheetnames[0]=='基线对比'
print('PASS: 原始7张表数据及73个链接保留；基线公式缓存、冻结窗格和首页顺序正确。')
